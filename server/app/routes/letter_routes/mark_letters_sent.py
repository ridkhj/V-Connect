from pathlib import Path

import openpyxl
from flask import Blueprint, current_app, jsonify, request


mark_letters_sent_bp = Blueprint('markletterssent', __name__, url_prefix='/mark-letters-sent')


@mark_letters_sent_bp.route('/<type>', methods=['POST'])
def mark_letters_sent(type):
    try:
        file_name, status_column = _resolve_type(type)

        payload = request.get_json(silent=True) or {}
        codes = payload.get('codes')

        if not isinstance(codes, list) or not all(isinstance(code, str) for code in codes):
            return jsonify({
                'success': False,
                'error': 'Erro de validacao',
                'details': 'O corpo da requisicao deve conter {"codes": ["COD1", "COD2"]}',
                'errorType': 'VALIDATION_ERROR'
            }), 400

        normalized_codes = {code.strip() for code in codes if code.strip()}
        if not normalized_codes:
            return jsonify({
                'success': False,
                'error': 'Erro de validacao',
                'details': 'A lista de codigos esta vazia',
                'errorType': 'VALIDATION_ERROR'
            }), 400

        sheets_path = Path(__file__).parent.parent.parent / 'assets' / 'sheets'
        file_path = sheets_path / file_name

        if not file_path.exists():
            return jsonify({
                'success': False,
                'error': 'Arquivo nao encontrado',
                'details': f'A planilha {file_name} nao foi encontrada. Processe os arquivos primeiro.',
                'errorType': 'FILE_NOT_FOUND'
            }), 404

        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active

        row = 2
        found_codes = set()

        while sheet.cell(row, 2).value is not None and sheet.cell(row, 3).value is not None:
            row_code = str(sheet.cell(row, 2).value).strip()
            if row_code in normalized_codes:
                sheet.cell(row, status_column).value = 'Enviado'
                found_codes.add(row_code)
            row += 1

        workbook.save(file_path)
        workbook.close()

        return jsonify({
            'success': True,
            'updated': len(found_codes),
            'message': 'Status atualizado com sucesso'
        }), 200

    except ValueError as e:
        return jsonify({
            'success': False,
            'error': 'Erro de validacao',
            'details': str(e),
            'errorType': 'VALIDATION_ERROR'
        }), 400
    except Exception as e:
        current_app.logger.error(f'Erro ao atualizar status de cartas: {str(e)}')
        return jsonify({
            'success': False,
            'error': 'Erro interno do servidor',
            'details': str(e),
            'errorType': 'SERVER_ERROR'
        }), 500


def _resolve_type(type):
    mapping = {
        'reciprocas': ('reciprocas.xlsx', 14),
        'nsl': ('nsl.xlsx', 12),
        'agradecimento': ('agradecimento.xlsx', 14)
    }

    if type not in mapping:
        raise ValueError('Tipo invalido. Tipos validos: reciprocas, nsl, agradecimento')

    return mapping[type]
