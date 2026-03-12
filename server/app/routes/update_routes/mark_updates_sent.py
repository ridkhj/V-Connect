from pathlib import Path

import openpyxl
from flask import Blueprint, current_app, jsonify, request


mark_updates_sent_bp = Blueprint('markupdatessent', __name__)


@mark_updates_sent_bp.route('/mark-updates-sent', methods=['POST'])
def mark_updates_sent():
    try:
        payload = request.get_json(silent=True) or {}
        codes = payload.get('codes')

        if not isinstance(codes, list) or not all(isinstance(code, str) for code in codes):
            return jsonify({
                'success': False,
                'error': 'Erro de validacao',
                'details': 'O corpo da requisicao deve conter {"codes": ["COD1", "COD2"]}',
                'errorType': 'VALIDATION_ERROR'
            }), 400

        normalized_codes = {code.strip() for code in codes if code.strip()}
        if not normalized_codes:
            return jsonify({
                'success': False,
                'error': 'Erro de validacao',
                'details': 'A lista de codigos esta vazia',
                'errorType': 'VALIDATION_ERROR'
            }), 400

        sheets_path = Path(__file__).parent.parent.parent / 'assets' / 'sheets'
        file_path = sheets_path / 'atualizacoes.xlsx'

        if not file_path.exists():
            return jsonify({
                'success': False,
                'error': 'Arquivo nao encontrado',
                'details': 'A planilha atualizacoes.xlsx nao foi encontrada. Processe os arquivos primeiro.',
                'errorType': 'FILE_NOT_FOUND'
            }), 404

        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active

        row = 2
        found_codes = set()

        while sheet.cell(row, 2).value is not None:
            row_code = str(sheet.cell(row, 2).value).strip()
            if row_code in normalized_codes:
                sheet.cell(row, 5).value = 'Enviado'
                found_codes.add(row_code)
            row += 1

        workbook.save(file_path)
        workbook.close()

        return jsonify({
            'success': True,
            'updated': len(found_codes),
            'message': 'Status atualizado com sucesso'
        }), 200

    except Exception as e:
        current_app.logger.error(f'Erro ao atualizar status de atualizacoes: {str(e)}')
        return jsonify({
            'success': False,
            'error': 'Erro interno do servidor',
            'details': str(e),
            'errorType': 'SERVER_ERROR'
        }), 500
